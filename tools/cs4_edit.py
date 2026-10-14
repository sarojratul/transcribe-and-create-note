#!/usr/bin/env python3
"""
cs4_edit.py: edit an existing 4-column cheat sheet (.docx made by cs4_post.py) in place, keeping
hand edits, numbering and styles. Any course. Topic blocks are paragraphs in style "Topic"; inside a
block the labels "Formulas", "Variables and constants", "Unit conversions", "Figures" mark sections.

  python3 cs4_edit.py dump   SHEET.docx
      List every topic block: formula / variable / figure counts.
  python3 cs4_edit.py append SHEET.docx FRAG.docx --topic "Speed and Speed Control" -o OUT.docx
      Add FRAG's Formulas items (numbering continues) and Variables items to the end of those lists.
      FRAG = pandoc + cs4_post.py output of a markdown file holding only **Formulas** / **Variables
      and constants** sections (use NEWX for added items, as on the main sheet).
  python3 cs4_edit.py addblock SHEET.docx FRAG.docx [--before "Colour Legend"] -o OUT.docx
      Insert FRAG's whole topic block(s) (Topic bar + sections) before a topic; lists restart at 1.
  python3 cs4_edit.py figs   SHEET.docx PLAN.yaml -o OUT.docx [--colw 2.45]
      PLAN: {"Topic name": [fig.png, "fig2.png@1.8", ...], ...}. Adds a "Figures" label at the end of
      each block (or reuses it) and the pictures under it. Width = true size from the PNG dpi
      (so all figures keep the same font size), capped at --colw; "path@W" forces W inches.
  python3 cs4_edit.py cut    SHEET.docx --topic "AC and Three-Phase Power" [--topic ...] -o OUT.docx
      Remove whole blocks (e.g. a midterm copy without later chapters).
  python3 cs4_edit.py legend SHEET.docx "Figures: navy arrows = current direction ..." -o OUT.docx
      Add a line to the Colour Legend just before its "Sizes:" line.
  python3 cs4_edit.py render SHEET.docx [--dir render] [--dpi 80]
      LibreOffice -> PDF (+ PNG per page), prints the page count. Installs libreoffice-math if
      missing. LibreOffice ignores equation size/colour, so it shows more pages than Word: compare
      against the original sheet rendered the same way, never quote the preview count as Word's.
Needs python-docx, lxml, pillow, pyyaml.
"""
import argparse, copy, os, shutil, subprocess, sys, tempfile
import docx
from docx.shared import Inches, Pt
from lxml import etree

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
w = lambda t: '{%s}%s' % (W, t)
LABELS = ('Formulas', 'Variables and constants', 'Unit conversions', 'Figures')


def txt(p):
    return ''.join(t.text or '' for t in p.iter(w('t'))).strip()


def style(p):
    ppr = p.find(w('pPr'))
    ps = ppr.find(w('pStyle')) if ppr is not None else None
    return ps.get(w('val')) if ps is not None else None


def numid(p):
    e = p.find('.//' + w('numId'))
    return e.get(w('val')) if e is not None else None


def setnum(p, v):
    p.find('.//' + w('numId')).set(w('val'), v)


def has_pic(p):
    return p.find('.//' + w('drawing')) is not None


def body_paras(body):
    return [p for p in body if p.tag == w('p')]


def blocks(body):
    """[(topic, [paragraphs incl. Topic bar])] in order."""
    out = []
    for p in body_paras(body):
        if style(p) == 'Topic':
            out.append((txt(p), [p]))
        elif out:
            out[-1][1].append(p)
    return out


def find_block(body, topic):
    for t, ps in blocks(body):
        if t == topic:
            return ps
    sys.exit('topic not found: %r (run dump)' % topic)


def sections(ps):
    s = {'f': [], 'v': [], 'u': [], 'g': [], 'labels': {}}
    cur = None
    key = {'Formulas': 'f', 'Variables and constants': 'v', 'Unit conversions': 'u', 'Figures': 'g'}
    for p in ps[1:]:
        t = txt(p)
        if t in key and not has_pic(p):
            cur = key[t]; s['labels'][cur] = p; continue
        if cur == 'g' and has_pic(p):
            s['g'].append(p)
        elif cur in ('f', 'v', 'u') and numid(p):
            s[cur].append(p)
    return s


def frag_body(path):
    d = docx.Document(path)
    return d, d.element.body


def cmd_dump(a):
    d = docx.Document(a.sheet)
    for t, ps in blocks(d.element.body):
        s = sections(ps)
        print('%-50s formulas %2d  variables %2d  conversions %d  figures %d' % (t, len(s['f']), len(s['v']), len(s['u']), len(s['g'])))


def cmd_append(a):
    d = docx.Document(a.sheet)
    tgt = sections(find_block(d.element.body, a.topic))
    _, fb = frag_body(a.frag)
    # fragment has no Topic bar: read its sections directly
    cur = None; ff = []; fv = []
    for p in body_paras(fb):
        t = txt(p)
        if t == 'Formulas': cur = 'f'; continue
        if t == 'Variables and constants': cur = 'v'; continue
        if numid(p):
            (ff if cur == 'f' else fv).append(p)
    for items, dest, name in ((ff, tgt['f'], 'Formulas'), (fv, tgt['v'], 'Variables and constants')):
        if not items: continue
        if not dest: sys.exit('block %r has no %s list to append to' % (a.topic, name))
        last = dest[-1]; nid = numid(last)
        for p in items:
            q = copy.deepcopy(p); setnum(q, nid); last.addnext(q); last = q
    d.save(a.o); print('appended %d formulas, %d variables to %r -> %s' % (len(ff), len(fv), a.topic, a.o))


def clone_num(numbering, src_id, new_id):
    root = numbering._element
    src = [n for n in root.findall(w('num')) if n.get(w('numId')) == src_id][0]
    c = copy.deepcopy(src); c.set(w('numId'), new_id)
    for k in list(c.attrib):
        if 'durableId' in k: del c.attrib[k]
    root.append(c)


def cmd_addblock(a):
    d = docx.Document(a.sheet); body = d.element.body
    anchor = find_block(body, a.before)[0]
    # templates: an existing numbered formula list and bullet variable list
    tf = tv = None
    for _, ps in blocks(body):
        s = sections(ps)
        if s['f'] and not tf: tf = numid(s['f'][0])
        if s['v'] and not tv: tv = numid(s['v'][0])
    numbering = d.part.numbering_part
    used = [int(n.get(w('numId'))) for n in numbering._element.findall(w('num'))]
    nxt = max(used) + 1
    fd, fb = frag_body(a.frag)
    # classify fragment lists by numFmt of level 0
    fnum = fd.part.numbering_part._element
    def is_bullet(nid):
        n = [x for x in fnum.findall(w('num')) if x.get(w('numId')) == nid][0]
        aid = n.find(w('abstractNumId')).get(w('val'))
        ab = [x for x in fnum.findall(w('abstractNum')) if x.get(w('abstractNumId')) == aid][0]
        return ab.find(w('lvl')).find(w('numFmt')).get(w('val')) == 'bullet'
    remap = {}
    for p in body_paras(fb):
        q = copy.deepcopy(p)
        nid = numid(q)
        if nid:
            if nid not in remap:
                remap[nid] = str(nxt); clone_num(numbering, tv if is_bullet(nid) else tf, str(nxt)); nxt += 1
            setnum(q, remap[nid])
        anchor.addprevious(q)
    d.save(a.o); print('inserted block(s) from %s before %r -> %s' % (a.frag, a.before, a.o))


def pic_width(path, colw):
    from PIL import Image
    if '@' in os.path.basename(path):
        path, wd = path.rsplit('@', 1); return path, float(wd)
    im = Image.open(path); dpi = (im.info.get('dpi') or (400, 400))[0] or 400
    return path, min(colw, im.size[0] / float(dpi))


def cmd_figs(a):
    import yaml
    plan = yaml.safe_load(open(a.plan))
    d = docx.Document(a.sheet); body = d.element.body
    n = 0
    for topic, figs in plan.items():
        ps = find_block(body, topic); s = sections(ps)
        if s['g']:
            last = s['g'][-1]
        elif 'g' in s['labels']:
            last = s['labels']['g']
        else:
            src = s['labels'].get('f') if s['labels'].get('f') is not None else s['labels'].get('v')
            lab = copy.deepcopy(src); ps[-1].addnext(lab)
            runs = lab.findall(w('r')); runs[0].find(w('t')).text = 'Figures'
            for r in runs[1:]: lab.remove(r)
            last = lab
        for f in figs:
            path, wd = pic_width(str(f), a.colw)
            p = d.add_paragraph(); pf = p.paragraph_format
            pf.space_before = Pt(1); pf.space_after = Pt(1); pf.keep_together = True; pf.line_spacing = 1.0
            p.add_run().add_picture(path, width=Inches(wd))
            last.addnext(p._p); last = p._p; n += 1
    d.save(a.o); print('inserted %d figures -> %s' % (n, a.o))


def cmd_cut(a):
    d = docx.Document(a.sheet); body = d.element.body
    for t in a.topic:
        for p in find_block(body, t): body.remove(p)
    d.save(a.o); print('removed %s -> %s' % (', '.join(a.topic), a.o))


def cmd_legend(a):
    d = docx.Document(a.sheet)
    sizes = [p for p in d.paragraphs if p.text.startswith('Sizes:')]
    if not sizes: sys.exit('no "Sizes:" legend line found')
    q = copy.deepcopy(sizes[0]._p); sizes[0]._p.addprevious(q)
    runs = q.findall(w('r')); runs[0].find(w('t')).text = a.text
    for r in runs[1:]: q.remove(r)
    d.save(a.o); print('legend line added -> %s' % a.o)


def cmd_render(a):
    if subprocess.run('dpkg -s libreoffice-math >/dev/null 2>&1', shell=True).returncode:
        subprocess.run('apt-get install -y -q libreoffice-math >/dev/null 2>&1 || (apt-get update -q >/dev/null 2>&1 && apt-get install -y -q libreoffice-math >/dev/null 2>&1)', shell=True)
    os.makedirs(a.dir, exist_ok=True)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', a.dir, a.sheet], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = os.path.join(a.dir, os.path.splitext(os.path.basename(a.sheet))[0] + '.pdf')
    pages = subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout
    print([l for l in pages.splitlines() if l.startswith('Pages')][0], '(LibreOffice preview, equations drawn larger than Word)')
    subprocess.run(['pdftoppm', '-r', str(a.dpi), '-png', pdf, os.path.join(a.dir, 'page')])
    print('PNGs in', a.dir)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest='cmd', required=True)
    p = sp.add_parser('dump'); p.add_argument('sheet'); p.set_defaults(f=cmd_dump)
    p = sp.add_parser('append'); p.add_argument('sheet'); p.add_argument('frag'); p.add_argument('--topic', required=True); p.add_argument('-o', required=True); p.set_defaults(f=cmd_append)
    p = sp.add_parser('addblock'); p.add_argument('sheet'); p.add_argument('frag'); p.add_argument('--before', default='Colour Legend'); p.add_argument('-o', required=True); p.set_defaults(f=cmd_addblock)
    p = sp.add_parser('figs'); p.add_argument('sheet'); p.add_argument('plan'); p.add_argument('--colw', type=float, default=2.45); p.add_argument('-o', required=True); p.set_defaults(f=cmd_figs)
    p = sp.add_parser('cut'); p.add_argument('sheet'); p.add_argument('--topic', action='append', required=True); p.add_argument('-o', required=True); p.set_defaults(f=cmd_cut)
    p = sp.add_parser('legend'); p.add_argument('sheet'); p.add_argument('text'); p.add_argument('-o', required=True); p.set_defaults(f=cmd_legend)
    p = sp.add_parser('render'); p.add_argument('sheet'); p.add_argument('--dir', default='render'); p.add_argument('--dpi', type=int, default=80); p.set_defaults(f=cmd_render)
    a = ap.parse_args(); a.f(a)


if __name__ == '__main__':
    main()
