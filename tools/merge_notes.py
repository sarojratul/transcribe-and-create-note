#!/usr/bin/env python3
"""
merge_notes.py: combine several lecture-notes .docx files (made by build_notes.py, or older
hand-made notes in the same family) into ONE Word file. Used by the document-merger skill.

  python3 merge_notes.py plan FILE_OR_DIR [...] -o plan.yaml [--lectures 4-18] [--title "..."]
  python3 merge_notes.py build plan.yaml -o OUT.docx

`plan` finds the notes files, reads each one's lecture number, date and topic from its file name
and title lines, sorts by lecture number, drops duplicates (same lecture: identical bytes, or the
older copy) and writes a YAML plan to check and edit. `build` merges in plan order:

  title: CME 331 Midterm Lecture Notes            # cover page title
  subtitle: Lectures 4, 6 and 9 to 18             # cover line (optional)
  footer: CME 331 · Midterm lecture notes         # one footer for the whole file, " · page N" added
  toc_levels: 1-2                                 # master contents depth
  max_dpi: 200                                    # downscale images above this (0 = keep as is)
  lectures:
    - file: /abs/path/CME331-lecture-04-notes.docx
      label: "Lecture 4, Sep 9: GPIO configuration and the init ritual"
  skipped:                                        # information only, never merged
    - {file: ..., reason: "same bytes as CME331-lecture-06-notes.docx"}

What build does to each lecture (content itself is never changed):
  - a Heading 1 divider with its label; its own headings move down one level (H1->H2, H2->H3,
    H3->H4), so the master contents reads lecture > section and slides stay in the nav pane;
  - its own "Contents" heading and TOC (field, SDT or static TOC lines, "press F9" hints) removed;
  - its page setup kept (portrait, landscape, 4-column cheat sheet): its last section's properties
    become a section break; page-number restarts and different-first-page flags removed;
  - per-lecture footers dropped; one generic footer with continuous page numbers.
Then: cover page + master TOC field (updateFields on, Word offers to fill it on open), styles of
the LAST lecture's file win (newest builder), schema fixes (numPr order), images downscaled to
max_dpi of their largest displayed width. Needs python-docx, lxml, pyyaml, pillow, docxcompose.
"""
import argparse, copy, filecmp, glob, io, math, os, re, sys, zipfile
import yaml
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
T = lambda e: "".join(t.text or "" for t in e.iter(qn("w:t")))
MONTHS = ("January February March April May June July August September October November "
          "December Jan Feb Mar Apr Jun Jul Aug Sep Sept Oct Nov Dec").split()
MON3 = {m: m[:3] for m in MONTHS}

# ---------------------------------------------------------------- plan
def first_lines(path, n=6):
    d = docx.Document(path)
    out = []
    for p in d.paragraphs:
        if p.text.strip():
            out.append(p.text.strip())
        if len(out) >= n:
            break
    return out

def detect(path):
    name = os.path.basename(path)
    lines = first_lines(path)
    head = " | ".join(lines)
    m = re.search(r"lec(?:ture)?[\s_-]*0*(\d+)", name, re.I) or re.search(r"Lecture\s+0*(\d+)", head)
    num = int(m.group(1)) if m else None
    date = ""
    md = re.search(r"\b(" + "|".join(MONTHS) + r")\.?\s+(\d{1,2})\b", head)
    if md:
        date = f"{MON3[md.group(1)]} {int(md.group(2))}"
    else:
        md = re.search(r"(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*[\s_-]*0?(\d{1,2})", name, re.I)
        if md:
            date = f"{md.group(1)[:3].title()} {int(md.group(2))}"
    topic = ""
    for ln in lines[1:]:
        mt = re.match(r"Topics?:\s*(.+)", ln)
        if mt:
            topic = mt.group(1); break
        mt = re.match(r"Lecture\s+\d+\s*(?:notes)?\s*:\s*(.+)", ln, re.I)
        if mt:
            topic = mt.group(1); break
    if not topic and len(lines) > 2 and not re.search(r"Instructor|Topics|\d{4}", lines[2]):
        topic = lines[2]
    topic = re.split(r";|\. ", topic)[0].strip().rstrip(".")
    topic = re.sub(r"\s*·\s*", ", ", topic)
    if len(topic) > 70:
        cut = topic[:70]
        topic = cut.rsplit(", ", 1)[0] if ", " in cut else cut.rsplit(" ", 1)[0]
    course = re.split(r"[:—|]", lines[0])[0].strip() if lines else ""
    label = f"Lecture {num}" if num is not None else os.path.splitext(name)[0]
    if date:
        label += f", {date}"
    if topic:
        label += f": {topic}"
    return dict(file=os.path.abspath(path), num=num, label=label, course=course,
                mtime=os.path.getmtime(path))

def parse_range(s):
    keep = set()
    for part in s.split(","):
        a, _, b = part.partition("-")
        keep |= set(range(int(a), int(b or a) + 1))
    return keep

def cmd_plan(a):
    files = []
    for p in a.inputs:
        if os.path.isdir(p):
            files += sorted(glob.glob(os.path.join(p, "**", "*.docx"), recursive=True))
        else:
            files.append(p)
    files = [f for f in files if not os.path.basename(f).startswith("~$")]
    info = [detect(f) for f in files]
    want = parse_range(a.lectures) if a.lectures else None
    by = {}
    skipped = []
    for i in info:
        if i["num"] is None:
            skipped.append({"file": i["file"], "reason": "no lecture number found"}); continue
        if want is not None and i["num"] not in want:
            skipped.append({"file": i["file"], "reason": f"lecture {i['num']} outside {a.lectures}"}); continue
        by.setdefault(i["num"], []).append(i)
    lectures = []
    for num in sorted(by):
        group = sorted(by[num], key=lambda x: -x["mtime"])
        keep = group[0]
        # prefer the standard name (lecture-NN-notes) when bytes are identical
        for g in group:
            if re.search(r"lecture-\d+-notes", os.path.basename(g["file"]), re.I) and filecmp.cmp(g["file"], keep["file"], shallow=False):
                keep = g
        for g in group:
            if g is keep:
                continue
            same = filecmp.cmp(g["file"], keep["file"], shallow=False)
            skipped.append({"file": g["file"], "reason": ("same bytes as " if same else "older or duplicate copy of lecture %d; kept " % num) + os.path.basename(keep["file"])})
        lectures.append({"file": keep["file"], "label": keep["label"]})
    course = next((i["course"] for i in info if i["course"]), "")
    code = re.match(r"[A-Z]{2,4}\s?\d{3}", course)
    nums = sorted(by)
    plan = {
        "title": a.title or f"{code.group(0) if code else course} Lecture Notes".strip(),
        "subtitle": f"Lectures {', '.join(map(str, nums))}" if nums else "",
        "footer": f"{code.group(0) if code else course} · Combined lecture notes".strip(" ·"),
        "toc_levels": "1-2",
        "max_dpi": 200,
        "lectures": lectures,
        "skipped": skipped,
    }
    with open(a.o, "w", encoding="utf-8") as f:
        yaml.safe_dump(plan, f, allow_unicode=True, sort_keys=False, width=1000)
    print(f"wrote {a.o}: {len(lectures)} lectures, {len(skipped)} skipped")
    for l in lectures:
        print("  ", l["label"], "<-", os.path.basename(l["file"]))
    for s in skipped:
        print("   skip", os.path.basename(s["file"]), ":", s["reason"])

# ---------------------------------------------------------------- build helpers
def pstyle(el):
    ps = el.find(".//" + qn("w:pStyle"))
    return ps.get(qn("w:val")) if ps is not None else ""

def clean_sectpr(sp, keep_refs=False):
    for tag in ("w:pgNumType", "w:titlePg"):
        for x in sp.findall(qn(tag)):
            sp.remove(x)
    if not keep_refs:
        for tag in ("w:headerReference", "w:footerReference"):
            for x in sp.findall(qn(tag)):
                sp.remove(x)
    return sp

def remove_contents(body):
    kids = list(body)
    for i, k in enumerate(kids):
        if k.tag == qn("w:p") and T(k).strip() in ("Contents", "Table of contents"):
            break
    else:
        return 0
    removed, depth = [k], 0
    for k in kids[i + 1:]:
        instr = "".join(x.text or "" for x in k.iter(qn("w:instrText")))
        fld = [f.get(qn("w:fldCharType")) for f in k.iter(qn("w:fldChar"))]
        txt = T(k)
        if k.tag == qn("w:sdt"):
            removed.append(k); continue
        if k.tag != qn("w:p"):
            break
        if pstyle(k).upper().startswith("TOC") or "TOC" in instr or "PAGEREF" in instr or depth > 0 or fld:
            depth += fld.count("begin") - fld.count("end")
            removed.append(k); continue
        if "F9" in txt or "looks blank" in txt or "list above" in txt:
            removed.append(k); continue
        break
    for k in removed:
        body.remove(k)
    return len(removed)

DEMOTE = {"Heading1": "Heading2", "Heading2": "Heading3", "Heading3": "Heading4"}

def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def heading_p(text, style="Heading1"):
    return parse_xml(f'<w:p xmlns:w="{W}"><w:pPr><w:pStyle w:val="{style}"/></w:pPr>'
                     f'<w:r><w:t xml:space="preserve">{esc(text)}</w:t></w:r></w:p>')

def prepare(path, label, last):
    d = docx.Document(path)
    body = d.element.body
    n = remove_contents(body)
    for ps in body.iter(qn("w:pStyle")):
        v = ps.get(qn("w:val"))
        if v in DEMOTE:
            ps.set(qn("w:val"), DEMOTE[v])
    body.insert(0, heading_p(label))
    for sp in body.iter(qn("w:sectPr")):
        clean_sectpr(sp)
    final = body.find(qn("w:sectPr"))
    if not last and final is not None:
        p = parse_xml(f'<w:p xmlns:w="{W}"><w:pPr/></w:p>')
        p.find(qn("w:pPr")).append(copy.deepcopy(final))
        final.addprevious(p)
    return d, n

def set_footer(master, sect, text):
    """Point the body section at one footer part holding `text · page N`."""
    fref = sect.find(qn("w:footerReference"))
    if fref is None:
        return None
    part = master.part.related_parts[fref.get(qn("r:id"))]
    root = part.element
    paras = root.findall(qn("w:p"))
    ppr = copy.deepcopy(paras[0].find(qn("w:pPr"))) if paras and paras[0].find(qn("w:pPr")) is not None else None
    r0 = paras[0].find(".//" + qn("w:r")) if paras else None
    rpr = copy.deepcopy(r0.find(qn("w:rPr"))) if r0 is not None and r0.find(qn("w:rPr")) is not None else None
    for p in list(root):
        root.remove(p)
    p = parse_xml(f'<w:p xmlns:w="{W}"/>')
    if ppr is not None:
        p.append(ppr)
    def run(inner):
        r = parse_xml(f'<w:r xmlns:w="{W}">{inner}</w:r>')
        if rpr is not None:
            r.insert(0, copy.deepcopy(rpr))
        p.append(r)
    run(f'<w:t xml:space="preserve">{esc(text)} · page </w:t>')
    run('<w:fldChar w:fldCharType="begin"/>')
    run('<w:instrText xml:space="preserve"> PAGE </w:instrText>')
    run('<w:fldChar w:fldCharType="separate"/>')
    run('<w:t>1</w:t>')
    run('<w:fldChar w:fldCharType="end"/>')
    root.append(p)
    # every other footer reference in the master points at the same part
    rid = fref.get(qn("r:id"))
    for x in sect.findall(qn("w:footerReference")):
        x.set(qn("r:id"), rid)
    for r_id, rel in list(master.part.rels.items()):
        if rel.reltype.endswith("/footer") and r_id != rid:
            master.part.drop_rel(r_id)
    return fref

SETTINGS_AFTER_UF = ["hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
                     "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
                     "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
                     "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
                     "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"]

def set_update_fields(settings):
    if settings.find(qn("w:updateFields")) is not None:
        return
    uf = etree.Element(qn("w:updateFields")); uf.set(qn("w:val"), "true")
    for c in settings:
        if etree.QName(c).localname in SETTINGS_AFTER_UF:
            c.addprevious(uf); return
    settings.append(uf)

NUMPR_AFTER = {"pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl"}

def fix_numpr(body):
    n = 0
    for ppr in body.iter(qn("w:pPr")):
        np_ = ppr.find(qn("w:numPr"))
        if np_ is None:
            continue
        lead = 0
        for c in ppr:
            if c is np_:
                break
            if etree.QName(c).localname in NUMPR_AFTER:
                lead += 1
            else:
                lead = -1; break
        if lead == -1:
            ppr.remove(np_)
            pos = sum(1 for c in ppr if etree.QName(c).localname in NUMPR_AFTER)
            ppr.insert(pos, np_); n += 1
    return n

def downscale(path, max_dpi):
    """Rewrite word/media images above max_dpi of their largest displayed width."""
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    from PIL import Image
    zin = zipfile.ZipFile(path)
    doc = etree.fromstring(zin.read("word/document.xml"))
    rels = etree.fromstring(zin.read("word/_rels/document.xml.rels"))
    rid2t = {r.get("Id"): r.get("Target") for r in rels}
    maxw = {}
    for dr in doc.iter(qn("w:drawing")):
        ext, blip = dr.find(".//" + WP + "extent"), dr.find(".//" + A + "blip")
        if ext is None or blip is None or blip.get(R + "embed") not in rid2t:
            continue
        key = "word/" + rid2t[blip.get(R + "embed")].lstrip("/").replace("word/", "", 1)
        maxw[key] = max(maxw.get(key, 0), int(ext.get("cx")) / 914400)
    tmp = path + ".tmp"
    zout = zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED)
    before = after = 0
    for info in zin.infolist():
        data = zin.read(info.filename)
        low = info.filename.lower()
        if info.filename in maxw and low.endswith((".png", ".jpg", ".jpeg")):
            before += len(data)
            im = Image.open(io.BytesIO(data))
            target = math.ceil(maxw[info.filename] * max_dpi)
            if im.width > target:
                im2 = im.resize((target, round(im.height * target / im.width)), Image.LANCZOS)
                buf = io.BytesIO()
                if low.endswith(".png"):
                    im2.save(buf, "PNG", optimize=True)
                else:
                    im2.convert("RGB").save(buf, "JPEG", quality=90)
                if len(buf.getvalue()) < len(data):
                    data = buf.getvalue()
            after += len(data)
        zout.writestr(info, data)
    zout.close(); zin.close()
    os.replace(tmp, path)
    return before, after

# ---------------------------------------------------------------- build
def cmd_build(a):
    from docxcompose.composer import Composer
    plan = yaml.safe_load(open(a.plan, encoding="utf-8"))
    lects = plan["lectures"]
    if not lects:
        sys.exit("plan has no lectures")
    for l in lects:
        if not os.path.exists(l["file"]):
            sys.exit("missing file: " + l["file"])
    if any("—" in str(plan.get(k, "")) for k in ("title", "subtitle", "footer")) or any("—" in l["label"] for l in lects):
        sys.exit("em dash in the plan text; rewrite it")

    master = docx.Document(lects[-1]["file"])
    mb = master.element.body
    msect = mb.find(qn("w:sectPr"))
    for k in list(mb):
        if k is not msect:
            mb.remove(k)
    clean_sectpr(msect, keep_refs=True)
    fref = set_footer(master, msect, plan.get("footer", ""))

    lv = plan.get("toc_levels", "1-2")
    cover = [
        parse_xml(f'<w:p xmlns:w="{W}"><w:pPr><w:spacing w:before="1200" w:after="120"/><w:jc w:val="center"/></w:pPr>'
                  f'<w:r><w:rPr><w:b/><w:sz w:val="48"/></w:rPr><w:t xml:space="preserve">{esc(plan.get("title", ""))}</w:t></w:r></w:p>'),
        parse_xml(f'<w:p xmlns:w="{W}"><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:rPr><w:sz w:val="24"/></w:rPr>'
                  f'<w:t xml:space="preserve">{esc(plan.get("subtitle", "") or "")}</w:t></w:r></w:p>'),
        parse_xml(f'<w:p xmlns:w="{W}"><w:pPr><w:spacing w:before="360"/></w:pPr><w:r><w:rPr><w:b/><w:sz w:val="28"/></w:rPr><w:t>Contents</w:t></w:r></w:p>'),
        parse_xml(f'<w:p xmlns:w="{W}"><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
                  f'<w:r><w:instrText xml:space="preserve"> TOC \\o "{lv}" \\h \\z \\u </w:instrText></w:r>'
                  f'<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
                  f'<w:r><w:t>Click Yes when Word asks to update fields, or right-click here and choose Update Field.</w:t></w:r>'
                  f'<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'),
    ]
    for it in cover:
        msect.addprevious(it)
    cover_sp = parse_xml(f'<w:sectPr xmlns:w="{W}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                         f'<w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" '
                         f'w:left="1440" w:header="720" w:footer="720" w:gutter="0"/><w:cols w:space="720"/></w:sectPr>')
    if fref is not None:
        cover_sp.insert(0, copy.deepcopy(fref))
    brk = parse_xml(f'<w:p xmlns:w="{W}"><w:pPr/></w:p>')
    brk.find(qn("w:pPr")).append(cover_sp)
    msect.addprevious(brk)

    comp = Composer(master)
    for i, l in enumerate(lects):
        d, n = prepare(l["file"], l["label"], last=(i == len(lects) - 1))
        comp.append(d)
        print(f"  + {l['label']}  ({os.path.basename(l['file'])}, removed {n} contents elements)")

    body = comp.doc.element.body
    last = docx.Document(lects[-1]["file"]).element.body.find(qn("w:sectPr"))
    bsp = body.find(qn("w:sectPr"))
    for tag in ("w:pgSz", "w:pgMar", "w:cols"):
        old, new = bsp.find(qn(tag)), last.find(qn(tag))
        if old is not None and new is not None:
            old.addprevious(copy.deepcopy(new)); bsp.remove(old)
    for sp in body.iter(qn("w:sectPr")):
        for tag in ("w:pgNumType", "w:titlePg"):
            for x in sp.findall(qn(tag)):
                sp.remove(x)
    moved = fix_numpr(body)
    set_update_fields(comp.doc.settings.element)
    comp.save(a.o)
    print(f"merged {len(lects)} files; numPr fixed: {moved}")
    if plan.get("max_dpi"):
        b, af = downscale(a.o, float(plan["max_dpi"]))
        print(f"images {b/1e6:.1f} MB -> {af/1e6:.1f} MB (max {plan['max_dpi']} dpi)")
    size = os.path.getsize(a.o) / 1e6
    print(f"saved {a.o}: {size:.1f} MB" + ("  (over 30 MB: too big for device_commit_files)" if size > 30 else ""))

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("plan"); p.add_argument("inputs", nargs="+"); p.add_argument("-o", required=True)
    p.add_argument("--lectures", help="keep only these lecture numbers, e.g. 1-18 or 4,6,9-18")
    p.add_argument("--title")
    b = sub.add_parser("build"); b.add_argument("plan"); b.add_argument("-o", required=True)
    a = ap.parse_args()
    (cmd_plan if a.cmd == "plan" else cmd_build)(a)

if __name__ == "__main__":
    main()
